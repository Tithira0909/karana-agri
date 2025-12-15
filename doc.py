from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_git_guide():
    doc = Document()

    # Title
    title = doc.add_heading('Essential Git Commands Guide', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('This guide covers the essential Git commands you need to master, organized by your daily workflow.')

    # 1. Setup
    doc.add_heading('1. First-Time Setup', level=1)
    doc.add_paragraph('Run these commands once when you install Git to configure your identity.')
    
    p = doc.add_paragraph()
    runner = p.add_run('git config --global user.name "Your Name"\ngit config --global user.email "you@example.com"')
    runner.font.name = 'Courier New'
    runner.font.color.rgb = RGBColor(0, 102, 204)

    # 2. Starting
    doc.add_heading('2. Starting a Project', level=1)
    doc.add_heading('Option A: Start from scratch', level=2)
    
    p = doc.add_paragraph()
    runner = p.add_run('cd my-project-folder\ngit init')
    runner.font.name = 'Courier New'
    runner.font.color.rgb = RGBColor(0, 102, 204)
    doc.add_paragraph('Creates a hidden .git folder. Your project is now tracked.', style='Quote')

    doc.add_heading('Option B: Download existing code', level=2)
    p = doc.add_paragraph()
    runner = p.add_run('git clone https://github.com/user/project.git')
    runner.font.name = 'Courier New'
    runner.font.color.rgb = RGBColor(0, 102, 204)

    # 3. Daily Workflow
    doc.add_heading('3. The Daily Workflow (The Save Loop)', level=1)
    doc.add_paragraph('Run this loop 90% of the time.')
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Step'
    hdr_cells[1].text = 'Command'
    hdr_cells[2].text = 'Description'

    data = [
        ('1. Check', 'git status', 'See changed/staged files.'),
        ('2. Stage', 'git add .', 'Move all changes to staging area.'),
        ('3. Save', 'git commit -m "msg"', 'Permanently save the snapshot.')
    ]

    for step, cmd, desc in data:
        row_cells = table.add_row().cells
        row_cells[0].text = step
        row_cells[1].text = cmd
        row_cells[2].text = desc

    # 4. Branching
    doc.add_heading('4. Branching (Parallel Work)', level=1)
    doc.add_paragraph('Use branches to isolate features from the main code.')
    p = doc.add_paragraph()
    runner = p.add_run('git checkout -b feature-login')
    runner.font.name = 'Courier New'
    runner.font.color.rgb = RGBColor(0, 102, 204)
    doc.add_paragraph('(Creates AND switches to a new branch)')

    # 5. Merging
    doc.add_heading('5. Merging (Combining Work)', level=1)
    p = doc.add_paragraph()
    runner = p.add_run('git checkout main\ngit pull origin main\ngit merge feature-login')
    runner.font.name = 'Courier New'
    runner.font.color.rgb = RGBColor(0, 102, 204)

    # 6. Undoing
    doc.add_heading('6. The "Oh No!" Commands', level=1)
    
    table2 = doc.add_table(rows=1, cols=2)
    table2.style = 'Table Grid'
    hdr_cells2 = table2.rows[0].cells
    hdr_cells2[0].text = 'Scenario'
    hdr_cells2[1].text = 'Command'

    undo_data = [
        ('Discard local file changes', 'git restore filename.py'),
        ('Unstage a file', 'git restore --staged filename.py'),
        ('Fix last commit message', 'git commit --amend'),
        ('Reset everything (Nuclear)', 'git reset --hard HEAD')
    ]

    for scenario, cmd in undo_data:
        row_cells = table2.add_row().cells
        row_cells[0].text = scenario
        row_cells[1].text = cmd

    # Cheat Sheet
    doc.add_page_break()
    doc.add_heading('CHEAT SHEET', level=0)
    doc.add_paragraph('1. New Work: git checkout -b branch-name')
    doc.add_paragraph('2. Stage: git add .')
    doc.add_paragraph('3. Save: git commit -m "message"')
    doc.add_paragraph('4. Upload: git push')

    doc.save('Git_Commands_Guide.docx')
    print("Word document created successfully!")

if __name__ == "__main__":
    create_git_guide()